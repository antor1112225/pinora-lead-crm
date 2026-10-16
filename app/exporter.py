"""
CSV and Excel Export Engine for Pinora-Like Local Lead Workspace.
Supports:
- Clean CSV generation with UTF-8 BOM encoding for seamless Excel compatibility
- Full XLSX export leveraging Python openpyxl (zero paid tools)
- Customizable column selection
- Proper escaping of commas, multiline notes, quotes
"""

import csv
import io
import datetime
from typing import List, Dict, Any, Optional
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

DEFAULT_EXPORT_COLUMNS = [
    ("business_name", "Business Name"),
    ("category", "Category"),
    ("specific_service", "Specific Service"),
    ("phone", "Public Phone"),
    ("email", "Public Email"),
    ("website", "Official Website"),
    ("address", "Address"),
    ("city", "City"),
    ("area", "Area"),
    ("postal_code", "Postal Code"),
    ("country", "Country"),
    ("crm_status", "CRM Status"),
    ("google_ads_status", "Google Ads Status"),
    ("meta_ads_status", "Meta Ads Status"),
    ("ads_summary", "Ads Verification Summary"),
    ("whatsapp_status", "WhatsApp Status"),
    ("next_action", "Next Scheduled Action"),
    ("next_followup_date", "Next Follow-up Date"),
    ("last_contact_date", "Last Contact Date"),
    ("last_contact_channel", "Last Contact Channel"),
    ("google_evidence_url", "Google Ads Evidence URL"),
    ("meta_evidence_url", "Meta Ads Evidence URL"),
    ("rating", "Rating"),
    ("reviews_count", "Reviews Count"),
    ("completeness_score", "Completeness Score (%)"),
    ("notes", "CRM Notes"),
    ("source_name", "Data Source"),
    ("research_date", "Research Date")
]


def generate_csv_bytes(leads: List[Dict[str, Any]], selected_fields: Optional[List[str]] = None) -> bytes:
    """Generate UTF-8 encoded CSV with BOM for universal Excel / Google Sheets compatibility."""
    output = io.StringIO()
    # Write BOM
    output.write('\ufeff')

    active_columns = []
    if selected_fields:
        for key, label in DEFAULT_EXPORT_COLUMNS:
            if key in selected_fields:
                active_columns.append((key, label))
    if not active_columns:
        active_columns = DEFAULT_EXPORT_COLUMNS

    writer = csv.writer(output, quoting=csv.QUOTE_MINIMAL)
    # Header row
    writer.writerow([label for _, label in active_columns])

    for lead in leads:
        row = []
        for key, _ in active_columns:
            val = lead.get(key, "")
            if val is None or val == "None":
                val = ""
            row.append(str(val))
        writer.writerow(row)

    return output.getvalue().encode('utf-8')


def generate_xlsx_bytes(leads: List[Dict[str, Any]], selected_fields: Optional[List[str]] = None) -> bytes:
    """Generate formatted Microsoft Excel XLSX workbook using openpyxl."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Local Leads"

    active_columns = []
    if selected_fields:
        for key, label in DEFAULT_EXPORT_COLUMNS:
            if key in selected_fields:
                active_columns.append((key, label))
    if not active_columns:
        active_columns = DEFAULT_EXPORT_COLUMNS

    # Header styling
    header_fill = PatternFill(start_color="0F1D2D", end_color="0F1D2D", fill_type="solid")
    header_font = Font(name="Segoe UI", size=11, bold=True, color="F8FAFC")
    regular_font = Font(name="Segoe UI", size=10, color="000000")
    border_thin = Border(
        left=Side(style='thin', color="E2E8F0"),
        right=Side(style='thin', color="E2E8F0"),
        top=Side(style='thin', color="E2E8F0"),
        bottom=Side(style='thin', color="E2E8F0")
    )

    # Write headers
    headers = [label for _, label in active_columns]
    ws.append(headers)

    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    # Write rows
    for row_idx, lead in enumerate(leads, start=2):
        row_vals = []
        for key, _ in active_columns:
            val = lead.get(key, "")
            if val is None or val == "None":
                val = ""
            row_vals.append(val)
        ws.append(row_vals)

        for col_idx in range(1, len(row_vals) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.font = regular_font
            cell.border = border_thin
            cell.alignment = Alignment(vertical="center")

    # Auto-adjust column widths
    for col in ws.columns:
        max_len = 0
        col_letter = col[0].column_letter
        for cell in col:
            val_str = str(cell.value or "")
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 40)

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()
